"""
Exportacion del reporte de calidad (Centro de Calidad IA) a Excel.

PARA QUE SIRVE
Llevar a gerencia, a la entidad o a una reunion de calibracion lo mismo que se
ve en pantalla, con el mismo filtro aplicado, en un archivo que se puede
auditar: cada numero del resumen se puede rastrear hasta sus filas.

DE DONDE SALEN LOS DATOS
Del navegador, ya filtrados. La pagina aplica los filtros (periodo, cartera,
supervisor...) sobre la reporteria y envia aqui las evaluaciones y los
criterios resultantes. Este modulo NO recalcula reglas de negocio: solo arma
el libro. Asi el Excel y la pantalla no pueden contar historias distintas.

COMO ESTA ORGANIZADO
  Resumen              periodo, filtros, meta y KPIs (formulas).
  Por cartera          nota, % en meta, error critico (formulas) y donde falla mas.
  Por supervisor       idem.
  Por agente           idem + tendencia y brecha que repite.
  Criterio x Cartera   % de cumplimiento por criterio (formulas) y N medidas.
  Criterio x Agente    idem por agente.
  Pareto               brechas por criterio con % acumulado (formulas).
  Evaluaciones         una fila por llamada (la base de todo lo anterior).
  Criterios            una fila por criterio evaluado en cada llamada.
  Definiciones         como se calcula cada indicador.

Los agregados son FORMULAS sobre las hojas de detalle, no numeros calculados
en Python: si alguien corrige una fila o cambia la meta en Resumen, todo se
recalcula. Lo que no es agregable con una formula simple (tendencia, brecha que
repite, "donde falla mas") viaja como valor ya calculado por la pagina y se
rotula como tal.

No inventa datos: lo que no viene en el payload queda vacio.
"""

from __future__ import annotations

from datetime import datetime, timedelta
from html import escape
from io import BytesIO
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.graphics.shapes import Circle, Drawing, Line, Rect, String
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FUENTE = "Arial"
AZUL_TITULO = "0D355B"
AZUL_CABECERA = "1F4E79"
GRIS_BORDE = "D5DEE8"
VERDE_FONDO, VERDE_TEXTO = "E3F4E9", "17633A"
AMBAR_FONDO, AMBAR_TEXTO = "FDF0D9", "8A5A0B"
ROJO_FONDO, ROJO_TEXTO = "FBE2E2", "A8323C"

BORDE = Border(bottom=Side(style="thin", color=GRIS_BORDE))
CABECERA_FILL = PatternFill("solid", fgColor=AZUL_CABECERA)
CABECERA_FONT = Font(name=FUENTE, bold=True, color="FFFFFF", size=10)
TEXTO = Font(name=FUENTE, size=10)
NEGRITA = Font(name=FUENTE, size=10, bold=True)
NOTA = Font(name=FUENTE, size=9, italic=True, color="5F7488")
INPUT_FONT = Font(name=FUENTE, size=11, bold=True, color="0000FF")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")

# Celda de la meta en Resumen: todas las formulas de "en meta" la referencian.
META_REF = "Resumen!$C$6"

COLS_EVAL = [
    ("id_feedback", "ID", 9),
    ("fecha", "Fecha llamada", 17),
    ("cartera", "Cartera", 26),
    ("agente", "Agente", 32),
    ("supervisor", "Supervisor", 28),
    ("tipo_llamada", "Tipo de llamada", 20),
    ("nota", "Nota vigente", 12),
    ("origen", "Origen de la nota", 14),
    ("nota_ia", "Nota IA", 10),
    ("evaluable", "Evaluable", 10),
    (None, "En meta", 9),  # formula
    ("error_critico", "Error crítico", 12),
    ("estado_revision", "Revisión supervisor", 16),
    ("ia_pide_revision", "IA pide revisión", 14),
    ("brechas", "Brechas (No cumple / Parcial)", 50),
    ("observacion", "Observación supervisor", 40),
]
COLS_CRIT = [
    ("id_feedback", "ID", 9),
    ("fecha", "Fecha llamada", 17),
    ("cartera", "Cartera", 26),
    ("agente", "Agente", 32),
    ("supervisor", "Supervisor", 28),
    ("tipo_llamada", "Tipo de llamada", 20),
    ("bloque", "Bloque", 9),
    ("bloque_nombre", "Nombre del bloque", 34),
    ("codigo", "Código", 10),
    ("criterio", "Criterio", 34),
    ("peso", "Peso", 7),
    ("nota", "Puntos", 8),
    ("resultado", "Resultado", 16),
    ("medido", "Medido (1/0)", 11),
    ("cumple", "Cumple (1/0)", 11),
    ("brecha", "Brecha (1/0)", 11),
]


def _col(cols, key: str) -> str:
    for i, (k, _h, _w) in enumerate(cols, start=1):
        if k == key:
            return get_column_letter(i)
    raise KeyError(key)


def _cabecera(ws: Worksheet, fila: int, titulos: List[str], anchos: Optional[List[int]] = None) -> None:
    for i, titulo in enumerate(titulos, start=1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.font = CABECERA_FONT
        c.fill = CABECERA_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if anchos:
            ws.column_dimensions[get_column_letter(i)].width = anchos[i - 1]
    ws.row_dimensions[fila].height = 30


def _titulo(ws: Worksheet, texto: str, subtitulo: str = "") -> None:
    ws["A1"] = texto
    ws["A1"].font = Font(name=FUENTE, size=14, bold=True, color=AZUL_TITULO)
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = NOTA


def _colores_pct(ws: Worksheet, rango: str) -> None:
    """Semaforo contra la meta de Resumen (verde >= meta, ambar >= 70, rojo)."""
    ws.conditional_formatting.add(rango, CellIsRule(operator="greaterThanOrEqual", formula=[f"{META_REF}/100"],
                                                    fill=PatternFill("solid", fgColor=VERDE_FONDO), font=Font(name=FUENTE, color=VERDE_TEXTO, bold=True)))
    ws.conditional_formatting.add(rango, CellIsRule(operator="between", formula=["0.7", f"{META_REF}/100-0.0001"],
                                                    fill=PatternFill("solid", fgColor=AMBAR_FONDO), font=Font(name=FUENTE, color=AMBAR_TEXTO, bold=True)))
    ws.conditional_formatting.add(rango, CellIsRule(operator="lessThan", formula=["0.7"],
                                                    fill=PatternFill("solid", fgColor=ROJO_FONDO), font=Font(name=FUENTE, color=ROJO_TEXTO, bold=True)))


def _escribir_detalle(ws: Worksheet, cols, filas: List[Dict], titulo: str, subtitulo: str) -> int:
    _titulo(ws, titulo, subtitulo)
    _cabecera(ws, 4, [h for _k, h, _w in cols], [w for _k, _h, w in cols])
    for r, fila in enumerate(filas, start=5):
        for c, (k, _h, _w) in enumerate(cols, start=1):
            if k is None:
                continue
            valor = fila.get(k)
            if k == "fecha" and isinstance(valor, str) and valor:
                try:
                    valor = datetime.fromisoformat(valor.replace("Z", "+00:00")).replace(tzinfo=None)
                except ValueError:
                    pass
            celda = ws.cell(row=r, column=c, value=valor if valor != "" else None)
            celda.font = TEXTO
            if k == "fecha" and isinstance(valor, datetime):
                celda.number_format = "dd/mm/yyyy hh:mm"
    ultima = max(5, 4 + len(filas))
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{ultima}"
    return ultima


def _rango(hoja: str, col: str, ultima: int) -> str:
    return f"'{hoja}'!${col}$5:${col}${ultima}"


def construir_reporte_calidad_excel(payload: Dict) -> bytes:
    evaluaciones: List[Dict] = payload.get("evaluaciones") or []
    criterios: List[Dict] = payload.get("criterios") or []
    meta = float(payload.get("meta") or 85)

    wb = Workbook()
    ws_res = wb.active
    ws_res.title = "Resumen"
    hojas = {n: wb.create_sheet(n) for n in ["Por cartera", "Por supervisor", "Por agente", "Criterio x Cartera",
                                             "Criterio x Agente", "Pareto", "Evaluaciones", "Criterios", "Definiciones"]}

    # --- detalle primero: todo lo demas lo referencia
    ult_e = _escribir_detalle(hojas["Evaluaciones"], COLS_EVAL, evaluaciones, "Evaluaciones",
                              "Una fila por llamada del periodo filtrado. Nota vacía = No evaluable.")
    ws_e = hojas["Evaluaciones"]
    c_nota, c_meta = _col(COLS_EVAL, "nota"), get_column_letter(11)
    for r in range(5, 5 + len(evaluaciones)):
        ws_e[f"{c_meta}{r}"] = f'=IF({c_nota}{r}="","",IF({c_nota}{r}>={META_REF},"Sí","No"))'
        ws_e[f"{c_meta}{r}"].font = TEXTO
        ws_e[f"{c_nota}{r}"].number_format = "0.0"
        ws_e[f"{_col(COLS_EVAL, 'nota_ia')}{r}"].number_format = "0.0"
        ws_e[f"{_col(COLS_EVAL, 'brechas')}{r}"].alignment = Alignment(wrap_text=True, vertical="top")

    ult_c = _escribir_detalle(hojas["Criterios"], COLS_CRIT, criterios, "Criterios por evaluación",
                              "Medido = Cumple, No cumple o Parcial. No aplica, No evaluable y Requiere revisión quedan con Medido = 0.")

    E = lambda key: _rango("Evaluaciones", _col(COLS_EVAL, key), ult_e)  # noqa: E731
    E_META = _rango("Evaluaciones", c_meta, ult_e)
    C = lambda key: _rango("Criterios", _col(COLS_CRIT, key), ult_c)  # noqa: E731

    # --- Resumen
    _titulo(ws_res, "Reporte de calidad — Centro de Calidad IA", f"Generado el {datetime.now():%d/%m/%Y %H:%M}"
            + (f" por {payload.get('generado_por')}" if payload.get("generado_por") else ""))
    ws_res.column_dimensions["A"].width = 3
    ws_res.column_dimensions["B"].width = 38
    ws_res.column_dimensions["C"].width = 26
    ws_res.column_dimensions["D"].width = 60
    ws_res["B4"], ws_res["C4"] = "Periodo", payload.get("periodo") or "Sin filtro de fecha"
    ws_res["B5"], ws_res["C5"] = "Filtros", "; ".join(f"{a}: {b}" for a, b in (payload.get("filtros") or [])) or "Ninguno"
    ws_res["B6"], ws_res["C6"] = "Meta de calidad (%)", meta
    ws_res["C6"].font, ws_res["C6"].fill = INPUT_FONT, INPUT_FILL
    ws_res["D6"] = "Editable: todas las columnas «en meta» y los colores usan esta celda."
    ws_res["D6"].font = NOTA
    for celda in ("B4", "B5", "B6"):
        ws_res[celda].font = NEGRITA
    _cabecera(ws_res, 8, ["", "Indicador", "Valor", "Cómo se calcula"])
    kpis = [
        ("Evaluaciones", f"=COUNTA({E('id_feedback')})", "0", "Llamadas del periodo filtrado."),
        ("Con nota", f"=COUNT({E('nota')})", "0", "Excluye las No evaluables."),
        ("No evaluables", "=C9-C10", "0", "Sin ningún criterio medible: fuera de los promedios."),
        ("Nota promedio", f"=IFERROR(AVERAGE({E('nota')}),\"\")", "0.0", "Promedio de la nota vigente (calibrada si Calidad la publicó)."),
        ("% en meta", f"=IFERROR(COUNTIF({E_META},\"Sí\")/C10,\"\")", "0.0%", "Evaluaciones con nota ≥ meta, sobre las que tienen nota."),
        ("Con error crítico", f"=COUNTIF({E('error_critico')},\"Sí\")", "0", "Error crítico, falta anulante o punto crítico."),
        ("Por revisar", f"=COUNTIF({E('estado_revision')},\"Por revisar\")", "0", "Sin cierre del supervisor."),
        ("Sin agente", f"=COUNTIF({E('agente')},\"Sin agente\")", "0", "No cuentan para ningún agente."),
        ("Criterios medidos (actos)", f"=SUM({C('medido')})", "0", "Criterio × llamada con resultado Cumple/No cumple/Parcial."),
        ("% de cumplimiento de criterios", f"=IFERROR(SUM({C('cumple')})/C17,\"\")", "0.0%", "Cumple / medidos."),
    ]
    for i, (nombre, formula, fmt, como) in enumerate(kpis, start=9):
        ws_res[f"B{i}"], ws_res[f"C{i}"], ws_res[f"D{i}"] = nombre, formula, como
        ws_res[f"B{i}"].font, ws_res[f"C{i}"].font, ws_res[f"D{i}"].font = NEGRITA, TEXTO, NOTA
        ws_res[f"C{i}"].number_format = fmt
        for col in "BCD":
            ws_res[f"{col}{i}"].border = BORDE
    ws_res["B21"] = "Las hojas «Por …» y «Criterio x …» se calculan con fórmulas sobre las hojas Evaluaciones y Criterios."
    ws_res["B21"].font = NOTA

    # --- Por cartera / supervisor / agente
    foco = payload.get("foco") or {}

    def hoja_grupo(ws: Worksheet, clave: str, etiqueta: str, extras: Optional[Dict[str, Dict]] = None) -> None:
        grupos = foco.get(clave) or []
        _titulo(ws, f"Por {etiqueta.lower()}", "Nota, % en meta y error crítico con fórmulas; «Dónde falla más» es el cálculo de la página (criterios con mayor % de falla).")
        titulos = [etiqueta, "Evaluaciones", "Con nota", "Nota promedio", "vs meta (pp)", "% en meta", "Con error crítico"]
        anchos = [34, 12, 10, 13, 12, 11, 13]
        if extras is not None:
            titulos += ["Tendencia (pp)", "Brecha que repite"]
            anchos += [13, 34]
        titulos += ["Dónde falla más (1)", "Dónde falla más (2)", "Dónde falla más (3)"]
        anchos += [36, 36, 36]
        _cabecera(ws, 4, titulos, anchos)
        col_e = _col(COLS_EVAL, clave)
        for r, g in enumerate(grupos, start=5):
            ws.cell(row=r, column=1, value=g.get("grupo")).font = NEGRITA
            ref = f"$A{r}"
            ws.cell(row=r, column=2, value=f"=COUNTIF({_rango('Evaluaciones', col_e, ult_e)},{ref})")
            ws.cell(row=r, column=3, value=f"=COUNTIFS({_rango('Evaluaciones', col_e, ult_e)},{ref},{E('nota')},\"<>\")")
            ws.cell(row=r, column=4, value=f"=IFERROR(AVERAGEIFS({E('nota')},{_rango('Evaluaciones', col_e, ult_e)},{ref}),\"\")").number_format = "0.0"
            ws.cell(row=r, column=5, value=f"=IF(D{r}=\"\",\"\",D{r}-{META_REF})").number_format = "+0.0;-0.0;0.0"
            ws.cell(row=r, column=6, value=f"=IFERROR(COUNTIFS({_rango('Evaluaciones', col_e, ult_e)},{ref},{E_META},\"Sí\")/C{r},\"\")").number_format = "0%"
            ws.cell(row=r, column=7, value=f"=COUNTIFS({_rango('Evaluaciones', col_e, ult_e)},{ref},{E('error_critico')},\"Sí\")")
            c = 8
            if extras is not None:
                ex = extras.get(g.get("grupo")) or {}
                t = ex.get("tendencia")
                ws.cell(row=r, column=c, value=t if t is not None else "sin base").number_format = "+0.0;-0.0;0.0"
                ws.cell(row=r, column=c + 1, value=ex.get("brecha_repite") or "")
                c += 2
            for j, txt in enumerate((g.get("top") or [])[:3]):
                ws.cell(row=r, column=c + j, value=txt).alignment = Alignment(wrap_text=True, vertical="top")
            for col in range(1, len(titulos) + 1):
                ws.cell(row=r, column=col).border = BORDE
                ws.cell(row=r, column=col).font = NEGRITA if col == 1 else TEXTO
        if grupos:
            ult = 4 + len(grupos)
            ws.conditional_formatting.add(f"D5:D{ult}", CellIsRule(operator="lessThan", formula=[META_REF], font=Font(name=FUENTE, color=ROJO_TEXTO, bold=True)))
            ws.auto_filter.ref = f"A4:{get_column_letter(len(titulos))}{ult}"
        ws.freeze_panes = "B5"

    hoja_grupo(hojas["Por cartera"], "cartera", "Cartera")
    hoja_grupo(hojas["Por supervisor"], "supervisor", "Supervisor")
    extras_ag = {a.get("agente"): a for a in (payload.get("agentes") or [])}
    hoja_grupo(hojas["Por agente"], "agente", "Agente", extras_ag)

    # --- matrices criterio x grupo
    criterios_catalogo = payload.get("catalogo_criterios") or []

    def matriz(ws: Worksheet, clave: str, etiqueta: str) -> None:
        grupos = [g.get("grupo") for g in (foco.get(clave) or [])]
        _titulo(ws, f"Cumplimiento por criterio y {etiqueta.lower()}",
                "% = llamadas que cumplen / llamadas donde el criterio se midió. Vacío = no se midió. Debajo, el N de medidas de cada celda.")
        _cabecera(ws, 4, ["Código", "Criterio"] + grupos + ["Total"], [10, 36] + [16] * len(grupos) + [12])
        col_g = _col(COLS_CRIT, clave)
        n = len(criterios_catalogo)
        fila_n = 4 + n + 3
        ws.cell(row=fila_n - 1, column=1, value="N de medidas por celda").font = NEGRITA
        _cabecera(ws, fila_n, ["Código", "Criterio"] + grupos + ["Total"])
        for i, crit in enumerate(criterios_catalogo):
            r, rn = 5 + i, fila_n + 1 + i
            for fila in (r, rn):
                ws.cell(row=fila, column=1, value=crit.get("codigo")).font = NEGRITA
                ws.cell(row=fila, column=2, value=crit.get("criterio")).font = TEXTO
            cond_cod = f"{C('codigo')},$A{r}"
            for j in range(len(grupos) + 1):
                col = get_column_letter(3 + j)
                if j < len(grupos):
                    cond = f"{cond_cod},{_rango('Criterios', col_g, ult_c)},{col}$4"
                else:
                    cond = cond_cod
                ws[f"{col}{r}"] = f"=IFERROR(COUNTIFS({cond},{C('cumple')},1)/COUNTIFS({cond},{C('medido')},1),\"\")"
                ws[f"{col}{r}"].number_format = "0%"
                ws[f"{col}{r}"].alignment = Alignment(horizontal="center")
                cond_n = cond.replace(f"$A{r}", f"$A{rn}").replace(f"{col}$4", f"{col}${fila_n}")
                ws[f"{col}{rn}"] = f"=COUNTIFS({cond_n},{C('medido')},1)"
                ws[f"{col}{rn}"].alignment = Alignment(horizontal="center")
                ws[f"{col}{rn}"].font = TEXTO
        if n:
            _colores_pct(ws, f"C5:{get_column_letter(3 + len(grupos))}{4 + n}")
        ws.freeze_panes = "C5"

    matriz(hojas["Criterio x Cartera"], "cartera", "Cartera")
    matriz(hojas["Criterio x Agente"], "agente", "Agente")

    # --- Pareto
    ws_p = hojas["Pareto"]
    _titulo(ws_p, "Pareto de brechas", "Brecha = No cumple o Parcial. Ordenado por número de brechas al momento de exportar.")
    _cabecera(ws_p, 4, ["Código", "Criterio", "Medidas", "Brechas", "% de falla", "% del total de brechas", "% acumulado"], [10, 36, 10, 10, 11, 16, 13])
    orden = payload.get("pareto") or []
    for i, crit in enumerate(orden):
        r = 5 + i
        ws_p[f"A{r}"], ws_p[f"B{r}"] = crit.get("codigo"), crit.get("criterio")
        ws_p[f"A{r}"].font = NEGRITA
        ws_p[f"C{r}"] = f"=COUNTIFS({C('codigo')},$A{r},{C('medido')},1)"
        ws_p[f"D{r}"] = f"=COUNTIFS({C('codigo')},$A{r},{C('brecha')},1)"
        ws_p[f"E{r}"] = f"=IFERROR(D{r}/C{r},\"\")"
        ws_p[f"F{r}"] = f"=IFERROR(D{r}/SUM($D$5:$D${4 + len(orden)}),\"\")"
        ws_p[f"G{r}"] = f"=IFERROR(SUM($D$5:D{r})/SUM($D$5:$D${4 + len(orden)}),\"\")"
        for col in "EFG":
            ws_p[f"{col}{r}"].number_format = "0.0%"
    ws_p.freeze_panes = "C5"

    # --- Definiciones
    ws_d = hojas["Definiciones"]
    _titulo(ws_d, "Definiciones")
    ws_d.column_dimensions["A"].width = 30
    ws_d.column_dimensions["B"].width = 110
    defs = [
        ("Nota vigente", "La nota de la llamada que cuenta: la calibrada si Calidad la publicó; si no, la de la IA."),
        ("No evaluable", "Llamada sin ningún criterio medible. Queda fuera de promedios y porcentajes y se cuenta aparte."),
        ("Criterio medido", "Resultado Cumple, No cumple o Parcial. No aplica, No evaluable y Requiere revisión no entran a ningún %."),
        ("Brecha", "Criterio medido con resultado No cumple o Parcial."),
        ("% de cumplimiento", "Llamadas donde el criterio cumple / llamadas donde se midió."),
        ("% en meta", "Evaluaciones con nota ≥ meta (Resumen!C6) / evaluaciones con nota."),
        ("Dónde falla más", "Los 3 criterios con mayor % de falla del grupo, con falla/medidas y la diferencia contra el % de falla de todo el periodo."),
        ("Tendencia", "Promedio de las 3 últimas notas del agente menos el de las 3 anteriores. «sin base» si no hay al menos 2 y 2."),
        ("Brecha que repite", "Regla acordada: el agente falla el mismo criterio 2+ veces en sus últimas 5 mediciones de ese criterio, con al menos 3 mediciones."),
        ("Base chica", "Menos de 3 medidas: el % existe pero no es concluyente."),
    ]
    for i, (a, b) in enumerate(defs, start=3):
        ws_d[f"A{i}"], ws_d[f"B{i}"] = a, b
        ws_d[f"A{i}"].font, ws_d[f"B{i}"].font = NEGRITA, TEXTO
        ws_d[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


# Los documentos ejecutivos no recalculan la evaluacion: reciben el mismo
# payload que la vista ya filtro. Asi respetan nombres, pesos y resultados de
# la pauta aplicada a cada llamada.
def _texto_documento(valor, vacio: str = "—") -> str:
    texto = str(valor or "").strip()
    return texto if texto else vacio


def _estilos_documento():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("quality-title", parent=base["Title"], fontName="Helvetica-Bold", fontSize=18,
                                leading=22, textColor=colors.HexColor("#10263D"), spaceAfter=3),
        "subtitle": ParagraphStyle("quality-subtitle", parent=base["Normal"], fontSize=8.5, leading=11,
                                   textColor=colors.HexColor("#5F7488"), spaceAfter=10),
        "section": ParagraphStyle("quality-section", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=10,
                                  leading=13, textColor=colors.HexColor("#10263D"), spaceBefore=10, spaceAfter=5),
        "body": ParagraphStyle("quality-body", parent=base["Normal"], fontSize=8, leading=10.5,
                               textColor=colors.HexColor("#23384D")),
        "small": ParagraphStyle("quality-small", parent=base["Normal"], fontSize=7, leading=9,
                                textColor=colors.HexColor("#415A73")),
        "header": ParagraphStyle("quality-header", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=7,
                                  leading=9, textColor=colors.white),
        "score": ParagraphStyle("quality-score", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=22,
                                leading=25, textColor=colors.HexColor("#10263D"), alignment=TA_CENTER),
    }


def _p(valor, estilo):
    return Paragraph(escape(_texto_documento(valor)), estilo)


def _p_lineas(valores, estilo, vacio: str = "—"):
    """Texto seguro con saltos reales para Paragraph; no deja HTML visible."""
    if isinstance(valores, (list, tuple)):
        partes = [str(x).strip() for x in valores if str(x).strip()]
    else:
        partes = [str(valores or "").strip()]
    return Paragraph("<br/>".join(escape(x) for x in partes) or escape(vacio), estilo)


def _fecha_documento(valor) -> str:
    texto = _texto_documento(valor)
    try:
        return datetime.fromisoformat(texto.replace("Z", "+00:00")).strftime("%d/%m/%Y %I:%M %p")
    except ValueError:
        return texto


def _semana_documento(valor) -> str:
    texto = _texto_documento(valor, "")
    try:
        fecha = datetime.fromisoformat(texto.replace("Z", "+00:00"))
        return f"Semana {fecha.isocalendar().week}"
    except ValueError:
        return "Sin fecha"


def _tabla_documento(datos, anchos, cabecera=True):
    tabla = Table(datos, colWidths=anchos, repeatRows=1 if cabecera else 0, hAlign="LEFT")
    comandos = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D5DEE8")),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if cabecera:
        comandos += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10263D")),
                     ("TEXTCOLOR", (0, 0), (-1, 0), colors.white)]
    tabla.setStyle(TableStyle(comandos))
    return tabla


def _tarjeta_documento(titulo: str, contenido, ancho: float, estilos, detalle: str = ""):
    encabezado = [_p(titulo, estilos["header"])]
    if detalle:
        encabezado.append(_p(detalle, estilos["header"]))
    card = Table([encabezado, [contenido]], colWidths=[ancho * .72, ancho * .28], hAlign="LEFT")
    card.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#10263D")),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#D5DEE8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("SPAN", (0, 1), (-1, 1)),
    ]))
    return card


def _fila_criterio_documento(fila: Dict, estilos):
    nombre = escape(_texto_documento(fila.get("criterio")))
    evidencia = _texto_documento(fila.get("evidencia"), "") if fila.get("brecha") else ""
    if evidencia:
        nombre += f'<br/><font size="6.5" color="#5F7488">{escape(evidencia)}</font>'
    nota, peso = fila.get("nota"), fila.get("peso")
    puntaje = "No medido" if nota is None else f"{nota} de {_texto_documento(peso, '0')}"
    resultado = _texto_documento(fila.get("resultado"))
    color = "#17633A" if resultado.lower() == "cumple" else "#A8323C" if "no cumple" in resultado.lower() else "#8A5A0B"
    return [Paragraph(nombre, estilos["small"]), _p(puntaje, estilos["small"]),
            Paragraph(f'<font color="{color}">{escape(resultado)}</font>', estilos["small"])]


def _bloques_criterios(criterios: List[Dict]) -> List[tuple[str, List[Dict]]]:
    bloques = {}
    for criterio in criterios or []:
        nombre = _texto_documento(criterio.get("bloque_nombre") or criterio.get("bloque"), "Criterios evaluados")
        bloques.setdefault(nombre, []).append(criterio)
    return list(bloques.items())


def construir_informe_individual_pdf(payload: Dict) -> bytes:
    """Informe horizontal de lectura ejecutiva y trazabilidad por criterio."""
    estilos = _estilos_documento()
    salida = BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=9 * mm, bottomMargin=9 * mm)
    criterios = payload.get("criterios") or []
    meta = payload.get("meta") or {}
    score = payload.get("score_vigente")
    score_txt = "No evaluable" if score is None else f"{float(score):.1f} / 100"

    cabecera = [
        Paragraph("Informe de evaluacion de llamada", estilos["title"]),
        Paragraph("Centro de Calidad | Resultado vigente de la llamada", estilos["subtitle"]),
        _tabla_documento([[
            _p_lineas(["AGENTE", _texto_documento(meta.get("agente"))], estilos["small"]),
            _p_lineas(["CARTERA", _texto_documento(meta.get("cartera"))], estilos["small"]),
            _p_lineas(["FECHA Y HORA", _fecha_documento(meta.get("fecha"))], estilos["small"]),
        ]], [64 * mm, 56 * mm, 53 * mm], cabecera=False),
    ]
    nota = _tarjeta_documento("NOTA FINAL DE LA LLAMADA", Paragraph(score_txt, estilos["score"]), 72 * mm, estilos, "Resultado vigente")
    historia = [Table([[cabecera, nota]], colWidths=[190 * mm, 77 * mm], style=[
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]), Spacer(1, 5),
        Paragraph("Resumen de la llamada", estilos["section"]),
        _tabla_documento([[_p(payload.get("resumen"), estilos["body"])]], [267 * mm], cabecera=False),
        Paragraph("Criterios de evaluacion", estilos["section"]),
    ]
    cards = []
    for bloque, filas in _bloques_criterios(criterios):
        suma_nota = sum(float(x.get("nota") or 0) for x in filas if x.get("nota") is not None)
        suma_peso = sum(float(x.get("peso") or 0) for x in filas)
        tabla = [_fila_criterio_documento(fila, estilos) for fila in filas]
        cards.append(_tarjeta_documento(bloque, _tabla_documento(tabla, [65 * mm, 20 * mm, 28 * mm], cabecera=False), 131 * mm, estilos, f"{suma_nota:g} de {suma_peso:g} pts"))
    for indice in range(0, len(cards), 2):
        fila = cards[indice:indice + 2]
        if len(fila) == 1:
            fila.append(Spacer(1, 1))
        historia.append(Table([fila], colWidths=[132 * mm, 132 * mm], style=[
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))

    fortalezas = payload.get("fortalezas") or []
    brechas = payload.get("brechas") or []
    historia += [
        Table([[
            _tarjeta_documento("FORTALEZAS OBSERVADAS", _p_lineas([f"• {x}" for x in fortalezas] or ["Sin fortaleza registrada."], estilos["body"]), 131 * mm, estilos),
            _tarjeta_documento("BRECHAS Y RECOMENDACION", _p_lineas([f"• {x}" for x in brechas] or ["Sin brecha registrada."], estilos["body"]), 131 * mm, estilos),
        ]], colWidths=[132 * mm, 132 * mm], style=[("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]),
        Table([[
            _tarjeta_documento("FEEDBACK DEL EVALUADOR", _p_lineas([
                _texto_documento(payload.get("feedback")), "",
                "Calidad | Firma: ____________________ | Fecha: ____________",
            ], estilos["body"]), 131 * mm, estilos),
            _tarjeta_documento("COMPROMISO DEL AGENTE", _p_lineas([
                _texto_documento(payload.get("compromiso")), "",
                "Jefatura de cartera (Lourdes) | Firma: ____________________ | Fecha: ____________",
                "Agente | Firma: ____________________ | Fecha: ____________",
            ], estilos["body"]), 131 * mm, estilos),
        ]], colWidths=[132 * mm, 132 * mm], style=[("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]),
    ]
    doc.build(historia)
    return salida.getvalue()


def _fecha_reporte(valor):
    texto = _texto_documento(valor, "")
    for formato in (None, "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.fromisoformat(texto.replace("Z", "+00:00")) if formato is None else datetime.strptime(texto[:10], formato)
        except ValueError:
            continue
    return None


def _barra_reporte(porcentaje: float, ancho: float = 31 * mm) -> Drawing:
    dibujo = Drawing(ancho, 7)
    dibujo.add(Rect(0, 1.5, ancho, 4, fillColor=colors.HexColor("#E3EAF2"), strokeColor=None))
    dibujo.add(Rect(0, 1.5, max(0, min(ancho, ancho * porcentaje / 100)), 4, fillColor=colors.HexColor("#1769C2"), strokeColor=None))
    return dibujo


def _grafico_evolucion(semanas: List[Dict]) -> Drawing:
    ancho, alto = 117 * mm, 30 * mm
    dibujo = Drawing(ancho, alto)
    izquierda, abajo, derecha, arriba = 18, 15, ancho - 5, alto - 8
    for valor in (60, 70, 80, 90, 100):
        y = abajo + (arriba - abajo) * (valor - 60) / 40
        dibujo.add(Line(izquierda, y, derecha, y, strokeColor=colors.HexColor("#E5ECF3"), strokeWidth=.4))
        dibujo.add(String(1, y - 2, str(valor), fontName="Helvetica", fontSize=5.5, fillColor=colors.HexColor("#5F7488")))
    dibujo.add(Line(izquierda, abajo, derecha, abajo, strokeColor=colors.HexColor("#7B8FA3"), strokeWidth=.5))
    if not semanas:
        dibujo.add(String(izquierda, alto / 2, "Sin semanas suficientes para mostrar evolución.", fontName="Helvetica", fontSize=7, fillColor=colors.HexColor("#5F7488")))
        return dibujo
    paso = (derecha - izquierda) / max(len(semanas) - 1, 1)
    puntos = []
    for indice, semana in enumerate(semanas):
        promedio = semana.get("promedio")
        if promedio is None:
            continue
        x = izquierda + paso * indice
        y = abajo + (arriba - abajo) * max(0, min(40, promedio - 60)) / 40
        puntos.append((x, y, promedio, semana.get("etiqueta", "")))
    for (x1, y1, _p1, _e1), (x2, y2, _p2, _e2) in zip(puntos, puntos[1:]):
        dibujo.add(Line(x1, y1, x2, y2, strokeColor=colors.HexColor("#1769C2"), strokeWidth=1.5))
    for x, y, promedio, etiqueta in puntos:
        dibujo.add(Circle(x, y, 2.1, fillColor=colors.HexColor("#1769C2"), strokeColor=colors.white, strokeWidth=.6))
        dibujo.add(String(x - 7, y + 5, f"{promedio:.1f}", fontName="Helvetica-Bold", fontSize=5.5, fillColor=colors.HexColor("#10263D")))
        dibujo.add(String(x - 9, 4, etiqueta, fontName="Helvetica", fontSize=5, fillColor=colors.HexColor("#415A73")))
    return dibujo


def construir_informe_cartera_pdf(payload: Dict) -> bytes:
    """Una pagina ejecutiva: tendencia, foco, agentes y plan desde datos reales."""
    estilos = _estilos_documento()
    salida = BytesIO()
    doc = SimpleDocTemplate(salida, pagesize=landscape(A4), leftMargin=7 * mm, rightMargin=7 * mm,
                            topMargin=7 * mm, bottomMargin=7 * mm)
    evaluaciones = payload.get("evaluaciones") or []
    criterios = payload.get("criterios") or []
    meta = float(payload.get("meta") or 80)
    evaluables = [x for x in evaluaciones if x.get("nota") is not None]
    promedio = sum(float(x["nota"]) for x in evaluables) / len(evaluables) if evaluables else None
    carteras = sorted({_texto_documento(x.get("cartera"), "") for x in evaluaciones if x.get("cartera")})
    cartera = carteras[0] if len(carteras) == 1 else "Carteras seleccionadas"
    criticos = sum(1 for x in evaluaciones if str(x.get("error_critico")).lower() in {"si", "sí", "true", "1"})

    por_agente = {}
    for fila in evaluaciones:
        por_agente.setdefault(_texto_documento(fila.get("agente"), "Sin agente"), []).append(fila)
    detalle_agentes = []
    for nombre, filas in por_agente.items():
        notas = [float(x["nota"]) for x in filas if x.get("nota") is not None]
        promedio_ag = sum(notas) / len(notas) if notas else None
        brecha = next((str(x.get("brecha_repite")) for x in payload.get("agentes", []) if x.get("agente") == nombre and x.get("brecha_repite")), "Sin recurrencia identificada")
        detalle_agentes.append({"nombre": nombre, "filas": filas, "n": len(filas), "promedio": promedio_ag, "brecha": brecha,
                                "criticos": sum(1 for x in filas if str(x.get("error_critico")).lower() in {"si", "sí", "true", "1"})})
    detalle_agentes.sort(key=lambda x: (x["promedio"] is None, x["promedio"] or 0))

    detalle_criterios = {}
    for fila in criterios:
        nombre = _texto_documento(fila.get("criterio"))
        registro = detalle_criterios.setdefault(nombre, {"nombre": nombre, "medidas": 0, "brechas": 0, "notas": []})
        if fila.get("medido"):
            registro["medidas"] += 1
            if fila.get("nota") is not None:
                registro["notas"].append(float(fila["nota"]))
        if fila.get("brecha"):
            registro["brechas"] += 1
    ranking = []
    for item in detalle_criterios.values():
        if item["medidas"]:
            item["porcentaje"] = item["brechas"] * 100 / item["medidas"]
            item["promedio"] = sum(item["notas"]) / len(item["notas"]) if item["notas"] else None
            ranking.append(item)
    ranking.sort(key=lambda x: (-x["porcentaje"], -x["brechas"], x["nombre"]))
    principal = ranking[0] if ranking else None
    medidos = sum(1 for x in criterios if x.get("medido"))
    cobertura = medidos * 100 / len(criterios) if criterios else None

    semanas_map = {}
    for fila in evaluaciones:
        fecha = _fecha_reporte(fila.get("fecha"))
        if not fecha:
            continue
        inicio = fecha.date() - timedelta(days=fecha.weekday())
        semanas_map.setdefault(inicio, []).append(fila)
    semanas = []
    for inicio, filas in sorted(semanas_map.items()):
        notas = [float(x["nota"]) for x in filas if x.get("nota") is not None]
        semanas.append({"etiqueta": f"S{len(semanas) + 1}", "promedio": sum(notas) / len(notas) if notas else None,
                        "n": len(filas), "criticos": sum(1 for x in filas if str(x.get("error_critico")).lower() in {"si", "sí", "true", "1"})})
    cambio = None
    if len(semanas) >= 2 and semanas[-1]["promedio"] is not None and semanas[-2]["promedio"] is not None:
        cambio = semanas[-1]["promedio"] - semanas[-2]["promedio"]

    encabezado = [
        Paragraph("Informe ejecutivo de calidad por cartera", estilos["title"]),
        Paragraph("MONITOREO DE LLAMADAS | CALIDAD | EXPERIENCIA DEL CLIENTE", estilos["subtitle"]),
        _tabla_documento([[
            _p_lineas(["CARTERA", cartera], estilos["small"]),
            _p_lineas(["PERIODO", _texto_documento(payload.get("periodo"))], estilos["small"]),
            _p_lineas(["MUESTRA EVALUADA", f"{len(evaluaciones)} llamadas"], estilos["small"]),
            _p_lineas(["FECHA DE EMISION", datetime.now().strftime("%d/%m/%Y")], estilos["small"]),
        ]], [65 * mm, 70 * mm, 57 * mm, 65 * mm], cabecera=False),
    ]
    def kpi(titulo, valor, detalle=""):
        contenido = f'<para align="center"><font size="18"><b>{escape(str(valor))}</b></font>'
        if detalle:
            contenido += f'<br/><font size="6.5" color="#5F7488">{escape(str(detalle))}</font>'
        contenido += "</para>"
        return _tarjeta_documento(titulo, Paragraph(contenido, estilos["body"]), 65 * mm, estilos)
    kpis = Table([[
        kpi("EVALUACIONES", str(len(evaluaciones)), f"{len(detalle_agentes)} agentes"),
        kpi("PROMEDIO DEL PERIODO", "—" if promedio is None else f"{promedio:.1f} / 100", "vs meta" if promedio is None else f"{promedio - meta:+.1f} pts vs meta"),
        kpi("RIESGO CRITICO", str(criticos), "evaluaciones con error critico"),
        kpi("COBERTURA MEDIDA", "—" if cobertura is None else f"{cobertura:.0f}%", "criterios con resultado medible"),
    ]], colWidths=[66 * mm, 66 * mm, 66 * mm, 66 * mm], style=[("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)])

    evolucion = _tarjeta_documento("EVOLUCION DE LA NOTA PROMEDIO", _grafico_evolucion(semanas), 131 * mm, estilos,
                                   "Comportamiento semanal del periodo")
    hito_lineas = []
    if cambio is not None:
        direccion = "mejora" if cambio >= 0 else "disminucion"
        hito_lineas.append(f"La nota promedio registra una {direccion} de {abs(cambio):.1f} puntos frente a la semana anterior.")
    elif promedio is not None:
        hito_lineas.append(f"La nota promedio del periodo es {promedio:.1f} sobre 100.")
    if principal:
        hito_lineas.append(f"Criterio mas afectado: {principal['nombre']} ({principal['porcentaje']:.0f}% de incumplimiento; {principal['brechas']} de {principal['medidas']} medidas).")
    hito_lineas.append(f"{criticos} evaluaciones presentan error critico." if criticos else "No se registran errores criticos en la muestra.")
    hitos = _tarjeta_documento("HITOS DEL PERIODO", _p_lineas([f"• {x}" for x in hito_lineas], estilos["body"]), 131 * mm, estilos,
                               "Principales logros y oportunidades")

    tabla_agentes = [[_p("#", estilos["header"]), _p("AGENTE", estilos["header"]), _p("EVAL.", estilos["header"]), _p("PROMEDIO", estilos["header"]), _p("PRINCIPAL BRECHA", estilos["header"]), _p("ESTADO", estilos["header"])]]
    for indice, fila in enumerate(detalle_agentes[:8], 1):
        estado = "En rango" if fila["promedio"] is not None and fila["promedio"] >= meta else "En seguimiento"
        tabla_agentes.append([_p(str(indice), estilos["small"]), _p(fila["nombre"], estilos["small"]), _p(str(fila["n"]), estilos["small"]),
                              _p("—" if fila["promedio"] is None else f"{fila['promedio']:.1f}", estilos["small"]), _p(fila["brecha"], estilos["small"]), _p(estado, estilos["small"])])
    agentes_card = _tarjeta_documento("RENDIMIENTO POR AGENTE", _tabla_documento(tabla_agentes, [8 * mm, 58 * mm, 17 * mm, 24 * mm, 101 * mm, 47 * mm]), 264 * mm, estilos,
                                      "Resultados individuales y oportunidad principal")

    tabla_criterios = [[_p("CRITERIO", estilos["header"]), _p("AFECTADAS", estilos["header"]), _p("%", estilos["header"]), _p("INCIDENCIA", estilos["header"])]]
    for fila in ranking[:5]:
        tabla_criterios.append([_p(fila["nombre"], estilos["small"]), _p(f"{fila['brechas']} de {fila['medidas']}", estilos["small"]),
                                _p(f"{fila['porcentaje']:.0f}%", estilos["small"]), _barra_reporte(fila["porcentaje"], 27 * mm)])
    criterios_card = _tarjeta_documento("CRITERIOS QUE REQUIEREN ATENCION", _tabla_documento(tabla_criterios, [57 * mm, 24 * mm, 13 * mm, 31 * mm]), 131 * mm, estilos,
                                        "Mayor incidencia de brechas")
    analisis = []
    if promedio is not None:
        analisis.append(f"La nota promedio es {promedio:.1f} sobre 100; la meta configurada es {meta:.0f}.")
    if principal:
        analisis.append(f"Priorizar {principal['nombre']} por concentrar la mayor incidencia de incumplimiento.")
    if detalle_agentes:
        prioridad = detalle_agentes[0]
        analisis.append(f"Foco inicial: {prioridad['nombre']} ({'sin nota evaluable' if prioridad['promedio'] is None else f'{prioridad['promedio']:.1f}'}).")
    analisis.append(_texto_documento(payload.get("recomendacion"), "Revisar semanalmente el avance por agente y formalizar compromisos de mejora."))
    analisis_card = _tarjeta_documento("ANALISIS Y RECOMENDACIONES", _p_lineas([f"{i + 1}. {texto}" for i, texto in enumerate(analisis)], estilos["body"]), 131 * mm, estilos,
                                      "Conclusiones y acciones sugeridas")
    plan = _tabla_documento([[
        _p("PLAN DE ACCION", estilos["header"]), _p("HALLAZGO PRIORITARIO", estilos["header"]), _p("RESPONSABLE", estilos["header"]), _p("SEGUIMIENTO", estilos["header"]),
    ], [
        _p("Coaching focalizado y seguimiento semanal.", estilos["small"]), _p(principal["nombre"] if principal else "Sin brechas medidas", estilos["small"]),
        _p("Calidad y jefatura de cartera", estilos["small"]), _p("Registrar avance en el plan de mejora.", estilos["small"]),
    ]], [43 * mm, 85 * mm, 62 * mm, 74 * mm])
    historia = encabezado + [Spacer(1, 3), kpis, Table([[evolucion, hitos]], colWidths=[132 * mm, 132 * mm], style=[("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]), agentes_card,
        Table([[criterios_card, analisis_card]], colWidths=[132 * mm, 132 * mm], style=[("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]), plan]
    doc.build(historia)
    return salida.getvalue()


def construir_informe_rtf(payload: Dict, tipo: str) -> bytes:
    """Formato editable que abre directamente en Word, sin requerir otra dependencia."""
    individual = tipo.upper() == "INDIVIDUAL"
    titulo = "INFORME DE EVALUACION DE LLAMADA" if individual else "INFORME EJECUTIVO POR CARTERA"
    lineas = [r"{\rtf1\ansi\deff0", r"{\fonttbl{\f0 Arial;}}", r"\fs24\b " + titulo + r"\b0\par"]
    if individual:
        meta = payload.get("meta") or {}
        lineas += [f"Nota final de la llamada: {_texto_documento(payload.get('score_vigente'), 'No evaluable')} / 100\\par",
                   f"Agente: {_texto_documento(meta.get('agente'))}\\par",
                   f"Cartera: {_texto_documento(meta.get('cartera'))}\\par",
                   f"Resumen: {_texto_documento(payload.get('resumen'))}\\par\\par",
                   r"\b CRITERIOS Y PUNTAJE OBTENIDO\b0\par"]
        for fila in payload.get("criterios") or []:
            nombre = _texto_documento(fila.get("criterio")).replace("\\", "\\\\")
            puntaje = "No medido" if fila.get("nota") is None else f"{fila.get('nota')} de {fila.get('peso', 0)}"
            lineas.append(f"{nombre}: {puntaje} | {_texto_documento(fila.get('resultado'))}\\par")
        lineas += [r"\par\b FEEDBACK DEL EVALUADOR\b0\par", _texto_documento(payload.get("feedback")) + r"\par",
                   r"\par\b COMPROMISO DEL AGENTE\b0\par", _texto_documento(payload.get("compromiso")) + r"\par\par",
                   "Calidad: ____________________    Jefatura de cartera (Lourdes): ____________________    Agente: ____________________\\par"]
    else:
        lineas += [f"Periodo: {_texto_documento(payload.get('periodo'))}\\par\\par", r"\b DETALLE POR AGENTE\b0\par"]
        for fila in payload.get("agentes") or []:
            lineas.append(f"{_texto_documento(fila.get('agente'))}: tendencia {_texto_documento(fila.get('tendencia'))}; brecha recurrente {_texto_documento(fila.get('brecha_repite'))}\\par")
        lineas += [r"\par\b RECOMENDACION EJECUTIVA\b0\par", _texto_documento(payload.get("recomendacion")) + r"\par"]
    lineas.append("}")
    return "\n".join(lineas).encode("latin-1", errors="replace")
